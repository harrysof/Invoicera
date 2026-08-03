"""
export_excel.py

Stage 3: takes a list of extract_fields.extract() results (one per
processed document) and writes them to a single .xlsx workbook with
two sheets:

  - "Factures"          -- documents where is_invoice == True
  - "Autres Documents"  -- everything else (Note d'Honoraires, parse
                            errors, unrecognized documents), using its
                            own column set rather than blank invoice
                            columns.

Rows flagged needs_review=True get a highlighted fill and their
review_reasons joined into a trailing column, so nothing gets silently
trusted -- you can filter/sort on that column in Excel directly.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from extract_fields import INVOICE_FIELDS, OTHER_DOC_FIELDS


HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
REVIEW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
BODY_FONT = Font(name="Arial")

INVOICE_COLUMNS = ["source_file", "document_type"] + INVOICE_FIELDS + ["needs_review", "review_reasons"]
OTHER_COLUMNS = ["source_file", "document_type"] + OTHER_DOC_FIELDS + ["needs_review", "review_reasons"]


def _write_sheet(ws, columns: list[str], rows: list[dict]) -> None:
    ws.append(columns)
    for col_idx, _ in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for row_data in rows:
        ws.append([row_data.get(c, "") for c in columns])
        row_idx = ws.max_row
        needs_review = row_data.get("needs_review")
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = BODY_FONT
            if needs_review:
                cell.fill = REVIEW_FILL

    # Reasonable auto-width, capped so review_reasons doesn't blow out the sheet
    for col_idx, col_name in enumerate(columns, start=1):
        max_len = max([len(col_name)] + [len(str(r.get(col_name, ""))) for r in rows]) if rows else len(col_name)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 45)

    ws.freeze_panes = "A2"


def build_workbook(results: list[dict], output_path: str) -> None:
    """
    results: list of dicts as returned by extract_fields.extract(), e.g.:
        {
            "source_file": "1640117789.pdf",
            "is_invoice": True,
            "document_type": "facture",
            "fields": {"RC": "...", "RC_source": "...", ...},
            "needs_review": False,
            "review_reasons": [],
        }
    """
    wb = Workbook()
    wb.remove(wb.active)

    invoice_rows = []
    other_rows = []

    for r in results:
        flat = {
            "source_file": r.get("source_file", ""),
            "document_type": r.get("document_type", ""),
            "needs_review": "YES" if r.get("needs_review") else "",
            "review_reasons": "; ".join(r.get("review_reasons", [])),
        }
        # Flatten fields, dropping the *_source keys from the sheet itself --
        # they're used for validation, not meant to clutter the deliverable.
        # (If you want them visible for audit, say so and I'll add a toggle.)
        for k, v in r.get("fields", {}).items():
            if not k.endswith("_source"):
                flat[k] = v if v is not None else ""

        if r.get("is_invoice"):
            invoice_rows.append(flat)
        else:
            other_rows.append(flat)

    ws_invoices = wb.create_sheet("Factures")
    _write_sheet(ws_invoices, INVOICE_COLUMNS, invoice_rows)

    ws_other = wb.create_sheet("Autres Documents")
    _write_sheet(ws_other, OTHER_COLUMNS, other_rows)

    wb.save(output_path)


if __name__ == "__main__":
    # Minimal smoke test with fabricated results (no live Ollama call).
    fake_results = [
        {
            "source_file": "1640117789.pdf",
            "is_invoice": True,
            "document_type": "facture",
            "fields": {
                "RC": "21 A 5910273-16/01", "RC_source": "N° RC : 21 A 5910273-16/01",
                "NIF": "19339110124317311601", "NIF_source": "NIF 19339110124317311601",
                "DATE": "25/05/2026", "DATE_source": "25/05/26",
                "FOURNISSEUR": "MAMME ABDELFATAH", "FOURNISSEUR_source": "MAMME ABDELFATAH",
                "N_FACTURE": "FC260118", "N_FACTURE_source": "FC260118",
                "PO_INVOICE": "1640117789", "PO_INVOICE_source": "1640117789",
                "MONTANT_HT": 19500.00, "MONTANT_HT_source": "19 500,00",
                "TVA": None, "TVA_source": None,
                "MONTANT_TTC": 19500.00, "MONTANT_TTC_source": "19 500,00",
            },
            "needs_review": False,
            "review_reasons": [],
        },
        {
            "source_file": "1640117822.pdf",
            "is_invoice": False,
            "document_type": "note_honoraires",
            "fields": {
                "DATE": "23/05/2026", "DATE_source": "23/05/2026",
                "EMETTEUR": "Pr. BENLAHRACH Zakia", "EMETTEUR_source": "Pr. BENLAHRACH Zakia",
                "DESTINATAIRE": "Roche Algerie SPA", "DESTINATAIRE_source": "Roche Algerie SPA",
                "REFERENCE": "CT2605AA6566", "REFERENCE_source": "CT2605AA6566",
                "MONTANT_BRUT": 164705.88, "MONTANT_BRUT_source": "164705.88",
                "RETENUE": 24705.88, "RETENUE_source": "24705.88",
                "NET_A_PAYER": 140000.00, "NET_A_PAYER_source": "140000.00",
                "PO_INVOICE": "1640117822", "PO_INVOICE_source": "1640117822",
            },
            "needs_review": False,
            "review_reasons": [],
        },
    ]
    build_workbook(fake_results, "/home/claude/invoicera_ext/smoke_test.xlsx")
    print("Wrote smoke_test.xlsx")
